"""
Ultra Genişletilmiş Trafik & Davranış Telemetrisi — Ek Analitikler ve Veri İhracı
Dosya: backend/telemetry_extras.py

- Cihaz / işletim sistemi / tarayıcı / ekran çözünürlüğü ayrıştırma
- Saatlik & günlük trafik yoğunluğu, giriş sayfaları, çıkış oranı (bounce), akış yolu
- Reel Getiri Simülatörü ve Finansal Sözlük kullanım istatistikleri
- Modül tıklanma sayıları ve fon bazlı "Reel Getiri testi" sayacı
- TÜBİTAK raporu için Excel (çok sayfalı) ve CSV (ham olay) dışa aktarımı
"""

import csv
import io
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

from database import get_connection

WEEKDAYS_TR = ["Pazar", "Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi"]
BENCHMARK_LABELS_TR = {
    "GOLD_GRAM_TRY": "Gram Altın",
    "GOLD_OUNCE_USD": "Ons Altın",
    "BIST100": "BIST 100",
    "USDTRY": "ABD Doları",
    "EURTRY": "Euro",
    "SP500": "S&P 500",
    "NASDAQ": "Nasdaq",
    "DEPOSIT": "Mevduat / TLREF",
}


def parse_user_agent(ua: str, screen_width: Optional[int] = None) -> Dict[str, str]:
    """Hafif (bağımlılıksız) User-Agent ayrıştırıcı."""
    s = ua or ""
    low = s.lower()

    if "iphone" in low or "ipod" in low:
        os_name, device = "iOS", "Mobil"
    elif "ipad" in low:
        os_name, device = "iOS", "Tablet"
    elif "android" in low:
        os_name = "Android"
        device = "Mobil" if "mobile" in low else "Tablet"
    elif "windows" in low:
        os_name, device = "Windows", "Masaüstü"
    elif "mac os x" in low or "macintosh" in low:
        os_name, device = "macOS", "Masaüstü"
    elif "cros" in low:
        os_name, device = "ChromeOS", "Masaüstü"
    elif "linux" in low:
        os_name, device = "Linux", "Masaüstü"
    else:
        os_name, device = "Diğer", "Bilinmiyor"

    if device in ("Bilinmiyor", "Masaüstü") and screen_width:
        if screen_width < 640:
            device = "Mobil"
        elif screen_width < 1024 and os_name in ("Diğer",):
            device = "Tablet"

    if "edg/" in low or "edge/" in low:
        browser = "Edge"
    elif "opr/" in low or "opera" in low:
        browser = "Opera"
    elif "samsungbrowser" in low:
        browser = "Samsung Internet"
    elif "firefox" in low or "fxios" in low:
        browser = "Firefox"
    elif "chrome" in low or "crios" in low:
        browser = "Chrome"
    elif "safari" in low:
        browser = "Safari"
    else:
        browser = "Diğer"

    return {"device_type": device, "os_name": os_name, "browser_name": browser}


def _safe_meta(raw: Any) -> Dict[str, Any]:
    if not raw:
        return {}
    try:
        data = json.loads(raw) if isinstance(raw, str) else dict(raw)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _pct(part: float, total: float) -> float:
    return round((part / total) * 100.0, 1) if total else 0.0


def _dist(counter: Counter, total: Optional[int] = None) -> List[Dict[str, Any]]:
    tot = total if total is not None else sum(counter.values())
    return [
        {"label": k, "count": v, "share_pct": _pct(v, tot)}
        for k, v in counter.most_common()
    ]


def compute_dashboard_extras(
    module_stats: List[Dict[str, Any]],
    fund_matrix: List[Dict[str, Any]],
    db_path: Optional[Path] = None,
) -> Dict[str, Any]:
    with get_connection(db_path) as conn:
        # ---------------- Trafik ----------------
        traffic_rows = conn.execute(
            """
            SELECT session_id, page_path, is_entry, device_type, os_name, browser_name,
                   screen_width, created_at
            FROM traffic_logs
            ORDER BY session_id, created_at, id
            """
        ).fetchall()

        sessions_pages: Dict[str, List[str]] = defaultdict(list)
        sess_device: Dict[str, str] = {}
        sess_os: Dict[str, str] = {}
        sess_browser: Dict[str, str] = {}
        sess_screen: Dict[str, int] = {}
        entry_counter: Counter = Counter()
        hour_counter: Counter = Counter()
        weekday_counter: Counter = Counter()
        page_counter: Counter = Counter()

        for r in traffic_rows:
            sid = r["session_id"]
            sessions_pages[sid].append(r["page_path"])
            sess_device[sid] = r["device_type"] or "Bilinmiyor"
            sess_os[sid] = r["os_name"] or "Diğer"
            sess_browser[sid] = r["browser_name"] or "Diğer"
            if r["screen_width"]:
                sess_screen[sid] = int(r["screen_width"])
            if r["is_entry"]:
                entry_counter[r["page_path"]] += 1
            page_counter[r["page_path"]] += 1
            try:
                dt = datetime.strptime(str(r["created_at"])[:19], "%Y-%m-%d %H:%M:%S")
                hour_counter[dt.hour] += 1
                weekday_counter[(dt.weekday() + 1) % 7] += 1  # Pazar=0
            except Exception:
                pass

        unique_visitors = len(sessions_pages)
        total_pageviews = len(traffic_rows)
        bounce_sessions = sum(1 for pages in sessions_pages.values() if len(pages) == 1)

        # Oturum süreleri (TIME_ON_PAGE toplamı)
        dur_rows = conn.execute(
            """
            SELECT session_id, SUM(duration_seconds) AS dur
            FROM analytics_events
            WHERE event_type = 'TIME_ON_PAGE'
            GROUP BY session_id
            """
        ).fetchall()
        dur_map = {r["session_id"]: float(r["dur"] or 0.0) for r in dur_rows}
        real_sess_durs = [dur_map[s] for s in sessions_pages if dur_map.get(s)]
        avg_session_sec = (sum(real_sess_durs) / len(real_sess_durs)) if real_sess_durs else 0.0

        # Anlık aktif kullanıcı (son 5 dk)
        cutoff = (datetime.now() - timedelta(minutes=5)).strftime("%Y-%m-%d %H:%M:%S")
        active_now = conn.execute(
            "SELECT COUNT(DISTINCT session_id) AS c FROM analytics_events WHERE created_at >= ?",
            (cutoff,),
        ).fetchone()["c"]

        # Akış yolu
        transitions: Counter = Counter()
        exit_counter: Counter = Counter()
        for pages in sessions_pages.values():
            for a, b in zip(pages, pages[1:]):
                if a != b:
                    transitions[f"{a} → {b}"] += 1
            exit_counter[pages[-1]] += 1

        screen_buckets: Counter = Counter()
        for w in sess_screen.values():
            if w < 640:
                screen_buckets["< 640 px (Telefon)"] += 1
            elif w < 1024:
                screen_buckets["640–1023 px (Tablet)"] += 1
            elif w < 1440:
                screen_buckets["1024–1439 px (Dizüstü)"] += 1
            else:
                screen_buckets["≥ 1440 px (Geniş Ekran)"] += 1

        hourly = [{"hour": h, "label": f"{h:02d}:00", "count": hour_counter.get(h, 0)} for h in range(24)]
        weekly = [
            {"weekday": i, "label": WEEKDAYS_TR[i], "count": weekday_counter.get(i, 0)}
            for i in (1, 2, 3, 4, 5, 6, 0)
        ]

        traffic_stats = {
            "unique_visitors": unique_visitors,
            "total_pageviews": total_pageviews,
            "active_now": int(active_now or 0),
            "avg_pages_per_session": round(total_pageviews / unique_visitors, 2) if unique_visitors else 0.0,
            "bounce_rate_pct": _pct(bounce_sessions, unique_visitors),
            "bounce_sessions": bounce_sessions,
            "avg_session_duration_seconds": round(avg_session_sec, 1),
            "avg_session_duration_minutes": round(avg_session_sec / 60.0, 2),
            "device_distribution": _dist(Counter(sess_device.values())),
            "os_distribution": _dist(Counter(sess_os.values())),
            "browser_distribution": _dist(Counter(sess_browser.values())),
            "screen_distribution": _dist(screen_buckets),
            "hourly_distribution": hourly,
            "weekday_distribution": weekly,
            "entry_pages": [{"page": k, "count": v} for k, v in entry_counter.most_common(8)],
            "exit_pages": [{"page": k, "count": v} for k, v in exit_counter.most_common(8)],
            "top_flows": [{"flow": k, "count": v} for k, v in transitions.most_common(10)],
            "top_pages": [{"page": k, "count": v} for k, v in page_counter.most_common(10)],
        }

        # ---------------- Modül tıklanma sayıları ----------------
        click_rows = conn.execute(
            """
            SELECT module_name, COUNT(*) AS c
            FROM analytics_events
            WHERE event_type NOT IN ('PAGE_VIEW', 'TIME_ON_PAGE')
            GROUP BY module_name
            """
        ).fetchall()
        clicks_map = {r["module_name"]: int(r["c"]) for r in click_rows}
        for m in module_stats:
            m["total_clicks"] = clicks_map.get(m["module_key"], 0)
            if m["module_key"] == "analytics":
                m["total_clicks"] += clicks_map.get("realreturn", 0)

        # ---------------- Reel Getiri Simülatörü ----------------
        rr_rows = conn.execute(
            """
            SELECT session_id, event_type, fund_code, metadata_json
            FROM analytics_events
            WHERE event_type IN ('REAL_RETURN_OPEN', 'REAL_RETURN_SIMULATE', 'REAL_RETURN_BENCHMARK_SELECT')
            """
        ).fetchall()
        rr_open_sessions, rr_sim_sessions = set(), set()
        open_cnt = sim_cnt = 0
        amounts: List[float] = []
        bench_counter: Counter = Counter()
        rr_fund_counter: Counter = Counter()
        for r in rr_rows:
            meta = _safe_meta(r["metadata_json"])
            if r["event_type"] == "REAL_RETURN_OPEN":
                open_cnt += 1
                rr_open_sessions.add(r["session_id"])
            elif r["event_type"] == "REAL_RETURN_SIMULATE":
                sim_cnt += 1
                rr_sim_sessions.add(r["session_id"])
                try:
                    if meta.get("initial_amount_tl") and not meta.get("_dup"):
                        amounts.append(float(meta["initial_amount_tl"]))
                except Exception:
                    pass
                if r["fund_code"]:
                    rr_fund_counter[r["fund_code"]] += 1
            else:
                key = str(meta.get("benchmark") or "")
                if key:
                    bench_counter[BENCHMARK_LABELS_TR.get(key, key)] += 1

        analysis_sessions = {
            r["session_id"]
            for r in conn.execute(
                """
                SELECT DISTINCT session_id FROM analytics_events
                WHERE event_type IN ('ANALYTICS_QUERY', 'FUND_COMPARE', 'BASKET_SIMULATE')
                """
            ).fetchall()
        }
        rr_open_of_analysis = len(rr_open_sessions & analysis_sessions) if analysis_sessions else len(rr_open_sessions)

        real_return_stats = {
            "open_clicks": open_cnt,
            "simulation_runs": sim_cnt,
            "unique_sessions_opened": len(rr_open_sessions),
            "analysis_sessions": len(analysis_sessions),
            "open_rate_pct": _pct(rr_open_of_analysis, len(analysis_sessions)),
            "avg_initial_amount_tl": round(sum(amounts) / len(amounts), 2) if amounts else 0.0,
            "benchmark_popularity": _dist(bench_counter),
            "top_funds": [{"fund_code": k, "count": v} for k, v in rr_fund_counter.most_common(10)],
        }

        # fon matrisi: reel getiri test sayısı
        fund_rr = {
            r["fund_code"]: int(r["c"])
            for r in conn.execute(
                """
                SELECT fund_code, COUNT(*) AS c FROM analytics_events
                WHERE event_type = 'REAL_RETURN_SIMULATE' AND fund_code IS NOT NULL
                GROUP BY fund_code
                """
            ).fetchall()
        }
        for f in fund_matrix:
            n = fund_rr.get(f["fund_code"], 0)
            f["real_return_tests"] = n
            f["total_inspections"] = int(f.get("total_inspections", 0)) + n
        fund_matrix.sort(
            key=lambda x: (x["total_inspections"], x["total_duration_seconds"]), reverse=True
        )

        # ---------------- Finansal Sözlük ----------------
        gl_rows = conn.execute(
            """
            SELECT event_type, metadata_json FROM analytics_events
            WHERE event_type IN ('GLOSSARY_TERM_READ', 'GLOSSARY_SEARCH')
            """
        ).fetchall()
        term_counter: Counter = Counter()
        cat_counter: Counter = Counter()
        search_counter: Counter = Counter()
        reads = searches = 0
        for r in gl_rows:
            meta = _safe_meta(r["metadata_json"])
            if r["event_type"] == "GLOSSARY_TERM_READ":
                reads += 1
                if meta.get("term"):
                    term_counter[str(meta["term"])] += 1
                if meta.get("category"):
                    cat_counter[str(meta["category"])] += 1
            else:
                searches += 1
                q = str(meta.get("query") or "").strip().lower()
                if q:
                    search_counter[q] += 1

        glossary_stats = {
            "total_term_reads": reads,
            "total_searches": searches,
            "top_terms": [{"term": k, "count": v} for k, v in term_counter.most_common(10)],
            "top_searches": [{"query": k, "count": v} for k, v in search_counter.most_common(10)],
            "category_breakdown": _dist(cat_counter),
        }

    return {
        "traffic_stats": traffic_stats,
        "real_return_stats": real_return_stats,
        "glossary_stats": glossary_stats,
    }


def build_events_csv(db_path: Optional[Path] = None) -> bytes:
    """Tüm ham telemetri olaylarını Excel uyumlu (UTF-8 BOM, ;) CSV olarak döndürür."""
    with get_connection(db_path) as conn:
        rows = conn.execute(
            """
            SELECT id, created_at, session_id, event_type, module_name, fund_code,
                   paired_fund_code, date_range_label, bulletin_id, duration_seconds,
                   basket_weight_pct, metadata_json
            FROM analytics_events ORDER BY id ASC
            """
        ).fetchall()
    out = io.StringIO()
    out.write("\ufeff")
    w = csv.writer(out, delimiter=";")
    w.writerow(
        [
            "ID", "Tarih", "Oturum", "Olay_Turu", "Modul", "Fon", "Eslesen_Fon",
            "Tarih_Araligi", "Bulten_ID", "Sure_Sn", "Sepet_Agirlik_Yuzde", "Ek_Veri_JSON",
        ]
    )
    for r in rows:
        w.writerow([r[k] if r[k] is not None else "" for k in r.keys()])
    return out.getvalue().encode("utf-8")


def build_academic_workbook(dashboard: Dict[str, Any], db_path: Optional[Path] = None) -> bytes:
    """Tüm istatistikleri çok sayfalı Excel (.xlsx) olarak TÜBİTAK raporu için dışa aktarır."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    def add_sheet(title: str, headers: List[str], rows: List[List[Any]], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 3, 12), 60)

    s = dashboard["summary"]
    add_sheet(
        "Genel Özet",
        ["Metrik", "Değer"],
        [
            ["Rapor Oluşturulma Zamanı", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Toplam Olay", s["total_events"]],
            ["Tekil Anonim Oturum", s["unique_sessions"]],
            ["Toplam Süre (dk)", s["total_duration_minutes"]],
            ["Toplam Fon İnceleme/Sorgu", s["total_fund_interactions"]],
            ["Toplam Karşılaştırma", s["total_comparisons"]],
            ["Toplam Sepet Simülasyonu", s["total_basket_simulations"]],
            ["Toplam Bülten Etkileşimi", s["total_bulletin_interactions"]],
        ],
        first=True,
    )

    t = dashboard["traffic_stats"]
    add_sheet(
        "Trafik",
        ["Metrik", "Değer"],
        [
            ["Tekil Ziyaretçi", t["unique_visitors"]],
            ["Toplam Sayfa Görüntüleme", t["total_pageviews"]],
            ["Anlık Aktif Kullanıcı", t["active_now"]],
            ["Oturum Başına Ort. Sayfa", t["avg_pages_per_session"]],
            ["Çıkış (Bounce) Oranı %", t["bounce_rate_pct"]],
            ["Ort. Oturum Süresi (sn)", t["avg_session_duration_seconds"]],
        ],
    )
    dist_rows: List[List[Any]] = []
    for name, key in (("Cihaz", "device_distribution"), ("İşletim Sistemi", "os_distribution"),
                      ("Tarayıcı", "browser_distribution"), ("Ekran", "screen_distribution")):
        for d in t[key]:
            dist_rows.append([name, d["label"], d["count"], d["share_pct"]])
    add_sheet("Cihaz-Tarayıcı-Ekran", ["Kategori", "Değer", "Oturum", "Pay %"], dist_rows)
    add_sheet("Saatlik Yoğunluk", ["Saat", "Sayfa Görüntüleme"], [[h["label"], h["count"]] for h in t["hourly_distribution"]])
    add_sheet("Günlük Yoğunluk", ["Gün", "Sayfa Görüntüleme"], [[d["label"], d["count"]] for d in t["weekday_distribution"]])
    add_sheet(
        "Akış Yolu",
        ["Tür", "Sayfa/Akış", "Adet"],
        [["Giriş", e["page"], e["count"]] for e in t["entry_pages"]]
        + [["Çıkış", e["page"], e["count"]] for e in t["exit_pages"]]
        + [["Akış", f["flow"], f["count"]] for f in t["top_flows"]],
    )
    add_sheet(
        "Modül İstatistikleri",
        ["Modül", "Yol", "Görüntülenme", "Tıklanma", "Tekil Oturum", "Toplam Süre (sn)", "Ort. Süre (sn)"],
        [[m["module_label"], m["route_path"], m["total_views"], m.get("total_clicks", 0),
          m["unique_sessions"], m["total_duration_seconds"], m["avg_duration_seconds"]] for m in dashboard["module_stats"]],
    )
    add_sheet(
        "Fon Etkileşim Matrisi",
        ["Fon", "Ünvan", "Toplam İnceleme", "Güven&Kanıt", "Şeffaflık", "Performans", "Reel Getiri Testi",
         "Sepete Ekleme", "Ort. Sepet Ağırlık %", "Toplam Süre (sn)", "Ort. Süre (sn)"],
        [[f["fund_code"], f["fund_title"], f["total_inspections"], f["evidence_views"], f["transparency_queries"],
          f["analytics_queries"], f.get("real_return_tests", 0), f["basket_additions"], f["avg_basket_weight_pct"],
          f["total_duration_seconds"], f["avg_duration_seconds"]] for f in dashboard["fund_matrix"]],
    )
    add_sheet(
        "Fon Çiftleri",
        ["Çift", "Karşılaştırma", "Pay %"],
        [[p["pair_label"], p["compare_count"], p["share_pct"]] for p in dashboard["comparison_stats"]["top_pairings"]],
    )
    add_sheet(
        "Vade Dağılımı",
        ["Vade", "Sorgu", "Pay %"],
        [[r["range_label"], r["query_count"], r["share_pct"]] for r in dashboard["comparison_stats"]["date_range_breakdown"]],
    )
    add_sheet(
        "Bülten İstatistikleri",
        ["Fon", "Başlık", "Kurum", "Açılma", "Kaynak Tıklama", "Ort. Süre (sn)"],
        [[b["fund_code"] or "GENEL", b["title"], b["institution"], b["open_count"], b["source_click_count"],
          b["avg_duration_seconds"]] for b in dashboard["bulletin_stats"]],
    )
    rr = dashboard["real_return_stats"]
    add_sheet(
        "Reel Getiri Simülatörü",
        ["Metrik", "Değer"],
        [
            ["Simülatör Açma Tıklaması", rr["open_clicks"]],
            ["Simülasyon Çalıştırma", rr["simulation_runs"]],
            ["Analiz Yapan Oturumların Açma Oranı %", rr["open_rate_pct"]],
            ["Ort. Başlangıç Bakiyesi (TL)", rr["avg_initial_amount_tl"]],
        ]
        + [[f"Benchmark: {b['label']}", b["count"]] for b in rr["benchmark_popularity"]],
    )
    gl = dashboard["glossary_stats"]
    add_sheet(
        "Finansal Sözlük",
        ["Tür", "Terim/Arama/Kategori", "Adet"],
        [["Okunan Terim", x["term"], x["count"]] for x in gl["top_terms"]]
        + [["Arama", x["query"], x["count"]] for x in gl["top_searches"]]
        + [["Kategori", x["label"], x["count"]] for x in gl["category_breakdown"]],
    )

    with get_connection(db_path) as conn:
        ev_rows = conn.execute(
            """
            SELECT id, created_at, session_id, event_type, module_name, fund_code, paired_fund_code,
                   date_range_label, bulletin_id, duration_seconds, basket_weight_pct
            FROM analytics_events ORDER BY id DESC LIMIT 20000
            """
        ).fetchall()
    add_sheet(
        "Ham Olaylar (Son 20K)",
        ["ID", "Tarih", "Oturum", "Olay", "Modül", "Fon", "Eşleşen Fon", "Vade", "Bülten ID", "Süre (sn)", "Sepet Ağırlık %"],
        [[r[k] if r[k] is not None else "" for k in r.keys()] for r in ev_rows],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
